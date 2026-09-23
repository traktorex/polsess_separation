#!/usr/bin/env python
"""Build the defence deck, thesis/my-writing/Prezentacja.pptx, from Prezentacja.md.

The markdown is the single source: one ``### N[A-Z]. Title`` heading per slide, and inside a section

    ![[file.png|width]]   the slide's figure (the Obsidian width is ignored; the file is looked up
                          under thesis/my-writing/figures/ and thesis/attachments/)
    - text                an on-slide bullet
    [anything]            ignored (the author's placeholders)
    other paragraphs      speaker notes

An empty title inherits the previous slide's title. Slides sharing a number (3A, 3B, 3C) form a
build sequence: their figures get one common scale and one common position, so nothing jumps when
the presenter steps through them. Slide 1 is the title slide, read from its ``Tytuł pracy:``,
``Imię i nazwisko:`` and ``promotor:`` lines plus the remaining lines (department, place and date).

Usage:
    venv/bin/python scripts/build_prezentacja.py            # writes the .pptx
    venv/bin/python scripts/build_prezentacja.py --export   # + PDF backup and PNG previews through
                                                            #   PowerPoint (Windows side, via PowerShell)
"""
from __future__ import annotations

import argparse
import io
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
VAULT = REPO / "thesis"
MD = VAULT / "my-writing" / "Prezentacja.md"
OUT = VAULT / "my-writing" / "Prezentacja.pptx"
FIG_DIRS = [VAULT / "my-writing" / "figures", VAULT / "attachments"]
LOGO = VAULT / "attachments" / "Pasted image 20260909175405.png"
WEBAPP_URL = "http://localhost:8871"          # the demo slide links here (webapp/run.sh default port)

# --- design tokens: the same ink and greys as the figures, DejaVu Sans everywhere ------------------
FONT = "DejaVu Sans"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x59, 0x63, 0x6E)
RULE = RGBColor(0xC3, 0xC2, 0xB7)

SLIDE_W, SLIDE_H = 13.333, 7.5                # inches, 16:9
MARGIN = 0.6
TITLE_TOP, TITLE_H, TITLE_PT = 0.40, 0.80, 28
LOGO_H = 0.40                                 # top-right corner of every slide
RULE_Y = 1.32                                 # hairline under the title zone
CONTENT_TOP, CONTENT_BOTTOM = 1.60, 6.90
BULLET_PT, SIDE_PT, BULLET_COL_W, WIDE_ASPECT = 22, 18, 4.0, 2.4   # bullets beside a figure, or below a wide one
NUM_PT = 11


@dataclass
class Slide:
    key: str                                  # "3A"
    num: int                                  # 3
    title: str
    figure: Path | None = None
    bullets: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)   # raw text lines (used by the title slide)


# --- markdown ----------------------------------------------------------------------------------------
HEADING = re.compile(r"^### (\d+)([A-Z]?)\.\s*(.*)$")
EMBED = re.compile(r"^!\[\[([^\]|]+)(?:\|[^\]]*)?\]\]\s*$")


def find_figure(name: str) -> Path:
    for d in FIG_DIRS:
        hits = sorted(d.rglob(name))
        if hits:
            return hits[0]
    raise FileNotFoundError(f"figure {name!r} not found under {[str(d) for d in FIG_DIRS]}")


def parse_md(path: Path) -> list[Slide]:
    slides: list[Slide] = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if m := HEADING.match(line):
            num, letter, title = int(m.group(1)), m.group(2), m.group(3).strip()
            if not title and slides:
                title = slides[-1].title
            slides.append(Slide(key=f"{num}{letter}", num=num, title=title))
            continue
        if not slides or not line.strip():
            continue
        s = slides[-1]
        if m := EMBED.match(line):
            s.figure = find_figure(m.group(1).strip())
        elif line.startswith("- "):
            s.bullets.append(line[2:].strip())
        elif line.startswith("[") and line.endswith("]"):
            continue
        else:
            s.notes.append(line.strip())
            s.lines.append(line.strip())
    return slides


def nbsp(text: str) -> str:
    """Polish typography for on-slide text: no line break after one-letter words, short
    abbreviations, or between a number and its unit."""
    text = re.sub(r"(?<![\w-])([aiouwzAIOUWZ]) ", "\\1\u00a0", text)
    text = re.sub(r"\b(np\.|ok\.|tj\.|dr|inż\.|mgr) ", "\\1\u00a0", text)
    text = re.sub(r"(\d) (s|h|min|kHz|dB)(?!\w)", "\\1\u00a0\\2", text)
    text = re.sub(r"(\d) %", "\\1\u00a0%", text)
    return text


# --- drawing helpers ---------------------------------------------------------------------------------
def add_text(slide, x, y, w, h, text, pt, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = nbsp(text)
    run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(pt), color
    return box


def add_bullets(slide, x, y, w, h, items, pt, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    hang = Inches(0.32)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(pt * 0.55)
        p.line_spacing = 1.1
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(hang))
        pPr.set("indent", str(-hang))
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        run = p.add_run()
        run.text = nbsp(item)
        run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(pt), INK
    return box


def logo_bytes() -> tuple[io.BytesIO, float]:
    """The PJATK logo trimmed to its ink; returns (png buffer, aspect ratio)."""
    im = Image.open(LOGO).convert("RGBA")
    alpha = im.getchannel("A").point(lambda a: 255 if a > 20 else 0)
    dark = im.convert("L").point(lambda v: 255 if v < 235 else 0)
    bbox = Image.composite(dark, Image.new("L", im.size, 0), alpha).getbbox()
    pad = 8
    trimmed = im.crop((bbox[0] - pad, bbox[1] - pad, bbox[2] + pad, bbox[3] + pad))
    buf = io.BytesIO()
    trimmed.save(buf, format="PNG")
    buf.seek(0)
    return buf, trimmed.size[0] / trimmed.size[1]


def chrome(slide, index: int, total: int, title: str | None, logo):
    """Logo top-right on every slide; title, hairline and slide number on content slides."""
    buf, aspect = logo
    buf.seek(0)
    logo_w = LOGO_H * aspect
    slide.shapes.add_picture(buf, Inches(SLIDE_W - MARGIN - logo_w), Inches(TITLE_TOP + 0.12),
                             height=Inches(LOGO_H))
    if title is None:
        return
    pt = TITLE_PT if len(title) <= 46 else TITLE_PT - 4      # keep every title on one line
    add_text(slide, MARGIN, TITLE_TOP, SLIDE_W - 2 * MARGIN - logo_w - 0.3, TITLE_H, title, pt,
             anchor=MSO_ANCHOR.MIDDLE)
    rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(MARGIN), Inches(RULE_Y),
                                      Inches(SLIDE_W - MARGIN), Inches(RULE_Y))
    rule.line.color.rgb = RULE
    rule.line.width = Pt(0.75)
    add_text(slide, SLIDE_W - MARGIN - 1.5, SLIDE_H - 0.50, 1.5, 0.3, f"{index} / {total}", NUM_PT,
             color=MUTED, align=PP_ALIGN.RIGHT)


# --- layout ------------------------------------------------------------------------------------------
def figure_box(s: Slide) -> tuple[float, float, float, float, str]:
    """(x, y, w, h, bullets_position) of the area the figure may fill."""
    x, y, w, h = MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
    if not s.bullets:
        return x, y, w, h, "none"
    px_w, px_h = Image.open(s.figure).size
    if px_w / px_h > WIDE_ASPECT:                       # wide figure: bullets underneath
        text_h = 0.55 * len(s.bullets) + 0.2
        return x, y, w, h - text_h - 0.15, "below"
    return x, y, w - BULLET_COL_W - 0.3, h, "right"     # otherwise a column on the right


def place_figures(prs, slides: list[Slide], pptx_slides):
    """Fit every figure into its box; slides sharing a number share scale and position."""
    groups: dict[int, list[Slide]] = {}
    for s in slides:
        if s.figure and s.num != 1:
            groups.setdefault(s.num, []).append(s)
    for members in groups.values():
        sizes = {s.key: Image.open(s.figure).size for s in members}
        boxes = {s.key: figure_box(s) for s in members}
        scale = min(min(b[2] / sizes[k][0], b[3] / sizes[k][1]) for k, b in boxes.items())
        widest = max(sizes[k][0] for k in sizes) * scale
        tallest = max(sizes[k][1] for k in sizes) * scale
        bx, by, bw = boxes[members[0].key][:3]
        bh = min(b[3] for b in boxes.values())
        left = bx + (bw - widest) / 2
        top = by + (bh - tallest) / 2
        for s in members:
            w, h = sizes[s.key][0] * scale, sizes[s.key][1] * scale
            pic = pptx_slides[s.key].shapes.add_picture(str(s.figure), Inches(left), Inches(top),
                                                        width=Inches(w), height=Inches(h))
            if s.num == 10:
                pic.click_action.hyperlink.address = WEBAPP_URL
            _, _, _, _, where = boxes[s.key]
            if where == "below":
                add_bullets(pptx_slides[s.key], MARGIN, top + h + 0.3, SLIDE_W - 2 * MARGIN,
                            CONTENT_BOTTOM - (top + h + 0.3), s.bullets, BULLET_PT - 2)
            elif where == "right":
                add_bullets(pptx_slides[s.key], SLIDE_W - MARGIN - BULLET_COL_W, CONTENT_TOP,
                            BULLET_COL_W, CONTENT_BOTTOM - CONTENT_TOP, s.bullets, SIDE_PT,
                            anchor=MSO_ANCHOR.MIDDLE)


def title_slide(slide, s: Slide, logo):
    chrome(slide, 0, 0, None, logo)
    title = author = None
    rest: list[str] = []
    for line in s.lines:
        low = line.lower()
        if low.startswith("tytuł pracy:"):
            title = line.split(":", 1)[1].strip()
        elif low.startswith("imię i nazwisko:"):
            author = line.split(":", 1)[1].strip()
        elif low.startswith("promotor:"):
            rest.append("Promotor: " + line.split(":", 1)[1].strip())
        else:
            rest.append(line)
    w = SLIDE_W - 2 * MARGIN
    add_text(slide, MARGIN, 1.9, w, 2.0, title or "", 28, anchor=MSO_ANCHOR.BOTTOM)
    add_text(slide, MARGIN, 4.25, w, 0.5, author or "", 22)
    y = 4.95
    for line in rest[:-1]:
        add_text(slide, MARGIN, y, w, 0.4, line, 18, color=MUTED)
        y += 0.42
    if rest:
        add_text(slide, MARGIN, SLIDE_H - 0.95, w, 0.4, rest[-1], 16, color=MUTED)   # place and date


def build(slides: list[Slide]) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    logo = logo_bytes()
    total = len(slides)
    pptx_slides = {}
    for i, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        pptx_slides[s.key] = slide
        if s.num == 1:
            title_slide(slide, s, logo)
        else:
            chrome(slide, i, total, s.title, logo)
            if s.bullets and not s.figure:
                add_bullets(slide, MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP,
                            s.bullets, BULLET_PT, anchor=MSO_ANCHOR.MIDDLE)
        if s.notes and s.num != 1:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(s.notes)
    place_figures(prs, slides, pptx_slides)
    prs.core_properties.title = "Obrona pracy magisterskiej"
    prs.core_properties.author = "Michał Dębski"
    return prs


# --- PowerPoint export (PDF backup + PNG previews) ---------------------------------------------------
def export_with_powerpoint(pptx: Path, preview_dir: Path) -> None:
    def win(p: Path) -> str:
        # resolve() turns the thesis/ symlink into its /mnt/<drive> target, so PowerPoint gets a
        # plain drive path instead of a \\wsl.localhost UNC path (SaveCopyAs rejects the latter)
        return subprocess.check_output(["wslpath", "-w", str(p.resolve())]).decode().strip()

    preview_dir.mkdir(parents=True, exist_ok=True)
    for old in preview_dir.glob("*.[pP][nN][gG]"):
        old.unlink()
    pdf = pptx.with_suffix(".pdf")
    pdf.unlink(missing_ok=True)
    # SaveAs embeds DejaVu Sans only when the target is a new file name, so save to a sibling and
    # swap it in afterwards. PowerPoint is single-instance: quit it only if this script started it.
    embedded = pptx.with_suffix(".embed.pptx")
    embedded.unlink(missing_ok=True)
    script = f"""
$ErrorActionPreference = "Stop"
$app = New-Object -ComObject PowerPoint.Application
$mine = ($app.Presentations.Count -eq 0)
$pres = $app.Presentations.Open("{win(pptx)}", 0, 0, 0)
try {{
    $pres.SaveAs("{win(embedded)}", 24, -1)      # ppSaveAsOpenXMLPresentation, EmbedTrueTypeFonts=msoTrue
    $pres.Export("{win(preview_dir)}", "PNG", 1920, 1080)
    $pres.SaveCopyAs("{win(pdf)}", 32)
}} finally {{
    $pres.Close()
    if ($mine) {{ $app.Quit() }}
}}
"""
    subprocess.run(["powershell.exe", "-NoProfile", "-Command", script], check=True)
    n_png = len(list(preview_dir.glob("*.[pP][nN][gG]")))
    if not (embedded.exists() and pdf.exists() and n_png):
        raise RuntimeError(f"PowerPoint export incomplete: embedded={embedded.exists()} "
                           f"pdf={pdf.exists()} previews={n_png}")
    embedded.replace(pptx)
    print(f"exported {pdf} (fonts embedded in {pptx.name}) and {n_png} previews to {preview_dir}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--export", action="store_true", help="PDF backup + PNG previews via PowerPoint")
    ap.add_argument("--preview-dir", type=Path, default=REPO / "docs" / "generated" / "prezentacja_preview")
    args = ap.parse_args()

    slides = parse_md(MD)
    prs = build(slides)
    prs.save(OUT)
    for i, s in enumerate(slides, 1):
        fig = s.figure.name if s.figure else "-"
        print(f"{i:2d}  {s.key:3s} {s.title[:48]:48s} {fig:45s} bullets={len(s.bullets)} notes={len(s.notes)}")
    print(f"wrote {OUT}")
    if args.export:
        export_with_powerpoint(OUT, args.preview_dir)


if __name__ == "__main__":
    main()
